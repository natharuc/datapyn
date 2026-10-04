"""Original data types, streaming formats, cooperative cancellation and SQL sessions."""

from decimal import Decimal
import io
import json
import os
from pathlib import Path
import sqlite3
import sys

import pandas as pd
import polars as pl
import pytest
from sqlalchemy import create_engine, text
from sqlalchemy.pool import QueuePool, StaticPool

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "source"))
from datapyn_runtime import data_tools
from datapyn_runtime.database import SQLiteConnector, ConnectorPool
from datapyn_runtime.export_control import ExportCancelled
from datapyn_runtime.export_formats import BoundedText
from datapyn_runtime.kernel import ResultStore
from datapyn_runtime.sql_export import value_literal
from test_runtime import Client


def export(frame, method="result.export_text", **params):
    return data_tools.dispatch(method, {"variable_name": "frame", **params}, {"frame": frame}, ResultStore(pd, pl))


@pytest.mark.parametrize("format", ["csv", "tsv", "txt", "excel"])
def test_clipboard_options_are_applied_to_the_original_sorted_selection(format):
    frame = pd.DataFrame({"id": [2**60, 2**60+1], "amount": [1.25, 2.5], "text": ["a\tb", 'a"b']})
    result = export(frame, format=format, options={"delimiter": "|", "decimal": ",", "include_header": False},
                    sort={"column": "amount", "direction": "desc"}, scope={"row_ranges": [[0, 0]]})
    delimiter = "\t" if format == "excel" else "|"
    rows = pd.read_csv(io.StringIO(result["text"]), sep=delimiter, header=None, decimal=",", dtype={0: str})
    assert rows.values.tolist() == [[str(2**60+1), 2.5, 'a"b']]
    assert result["row_count"] == 1 and result["column_count"] == 3


@pytest.mark.parametrize("orient", ["records", "split", "index", "columns", "values", "table"])
def test_json_orientations_are_valid_and_preserve_values(tmp_path, orient):
    frame = pd.DataFrame({"value": [3, 7], "text": ["ação", None]}, index=pd.Index([5, 9], name="row"))
    content = export(frame, format="json", options={"orient": orient, "index": True, "indent": 2})["text"]
    parsed = json.loads(content)
    loaded = pd.read_json(io.StringIO(content), orient=orient)
    assert loaded["value"].tolist() == [3, 7] if orient != "values" else loaded.iloc[:, 0].tolist() == [3, 7]
    assert "ação" in content
    if orient in {"split", "table", "index", "columns"}:
        assert loaded.index.tolist() == [5, 9]


def test_json_lines_are_one_compact_record_per_line_despite_indent():
    content = export(pd.DataFrame({"a": [1, 2]}), format="json", options={"orient": "records", "lines": True, "indent": 4})["text"]
    assert [json.loads(line) for line in content.splitlines()] == [{"a": 1}, {"a": 2}]
    with pytest.raises(ValueError, match="requires records"):
        export(pd.DataFrame({"a": [1]}), format="json", options={"orient": "split", "lines": True})


def test_clipboard_size_guard_and_precision():
    frame = pd.DataFrame({"a": [Decimal("1.234567890123456789"), 2**63, b"\x00\x7f"]})
    parsed = json.loads(export(frame, format="json")["text"])
    assert parsed == [{"a": "1.234567890123456789"}, {"a": str(2**63)}, {"a": "007f"}]
    with pytest.raises(ValueError, match="16 MiB"):
        export(pd.DataFrame({"value": ["x" * (16 * 1024 * 1024)]}), format="csv")
    output = BoundedText(8)
    output.write("ação")
    with pytest.raises(ValueError):
        output.write("ção")


def test_xlsx_custom_sheet_index_headers_and_formula_literal(tmp_path):
    from openpyxl import load_workbook
    frame = pd.DataFrame({"formula": ["=1+1"], "date": [pd.NaT]}, index=pd.Index([12], name="row"))
    path = tmp_path / "data.xlsx"
    export(frame, "result.export", path=str(path), options={"sheet_name": "Dados", "index": True, "include_header": False})
    book = load_workbook(path)
    assert book.sheetnames == ["Dados"]
    assert book.active["A1"].value == 12 and book.active["C1"].value is None
    assert book.active["B1"].data_type == "s"
    book.close()


def test_parquet_chunks_keep_schema_when_first_group_is_all_null(tmp_path):
    frame = pd.DataFrame({"value": pd.Series([None] * 100000 + ["last"], dtype=object)})
    path = tmp_path / "data.parquet"
    export(frame, "result.export", path=str(path), options={"compression": "zstd"})
    pd.testing.assert_frame_equal(pd.read_parquet(path), frame)


def test_sql_create_insert_executes_with_original_bigint_binary_and_escaping():
    frame = pd.DataFrame({"big": [2**60+1], "text": ["O'Reilly\\path"], "blob": [b"\x00\xff"], "decimal": [Decimal("12.345")], "flag": [True]})
    content = export(frame, format="sql", options={"db_type": "sqlite", "table_name": 'main."a.b"', "sql_mode": "create_insert", "batch_size": 1000, "include_transaction": True})["text"]
    with sqlite3.connect(":memory:") as connection:
        connection.executescript(content)
        assert connection.execute('SELECT * FROM "a.b"').fetchone() == (2**60+1, "O'Reilly\\path", b"\x00\xff", "12.345", 1)
    assert value_literal(Decimal("12.345"), "postgresql") == "12.345"
    assert value_literal(b"\x00\xff", "sqlserver") == "0x00ff"
    assert value_literal(b"\x00\xff", "postgresql") == "decode('00ff', 'hex')"
    assert value_literal("ação\\p'", "mysql") == "'ação\\\\p'''"
    assert value_literal("ação", "sqlserver") == "N'ação'"


def test_sqlserver_insert_batches_obey_server_limit_and_go():
    content = export(pd.DataFrame({"value": range(1001)}), format="sql", options={"batch_size": 100000, "table_name": "[dbo].[a]]b]", "include_go": True})["text"]
    assert content.count("INSERT INTO [dbo].[a]]b]") == 2 and content.endswith("GO\n")
    with pytest.raises(ValueError, match="cannot include CREATE"):
        export(pd.DataFrame({"x": [1]}), format="sql", options={"db_type": "databricks", "sql_mode": "create_insert", "include_transaction": True})
    content = export(pd.DataFrame({"x": [1]}), format="sql", options={"db_type": "databricks", "include_transaction": True})["text"]
    assert "BEGIN TRANSACTION;" in content


@pytest.mark.parametrize("format", ["csv", "xlsx", "json", "parquet", "sql"])
def test_cancelled_file_export_keeps_existing_destination_and_cleans_work(tmp_path, format):
    frame = pd.DataFrame({"a": range(1100)})
    path = tmp_path / f"keep.{format}"
    path.write_bytes(b"existing")
    cancelled, progress = [False], []
    def update(item):
        progress.append(item)
        if item["phase"] == "writing":
            cancelled[0] = True
    with pytest.raises(ExportCancelled):
        data_tools.dispatch("result.export", {"variable_name": "frame", "path": str(path)}, {"frame": frame}, ResultStore(pd, pl),
                            progress=update, cancelled=lambda: cancelled[0])
    assert path.read_bytes() == b"existing" and list(tmp_path.iterdir()) == [path]
    assert progress[-1]["phase"] == "cancelled"


def test_temporary_sqlite_table_preserves_permanent_namesake_and_survives_checkouts(tmp_path):
    connector = SQLiteConnector(str(tmp_path / "temp.sqlite"))
    try:
        connector.execute_query('CREATE TABLE shadow(a INTEGER); INSERT INTO shadow VALUES(99);')
        frame = pd.DataFrame({"a": [1, 2]})
        namespace, store = {"frame": frame}, ResultStore(pd, pl)
        params = {"variable_name": "frame", "table": "shadow", "temporary": True}
        result = data_tools.dispatch("result.export_table", params, namespace, store, connector)
        assert result["temporary"] and connector.has_temporary_tables
        assert connector.execute_query("SELECT * FROM shadow")["a"].tolist() == [1, 2]
        assert connector.execute_query("SELECT * FROM main.shadow")["a"].tolist() == [99]
        for mode, expected in [("append", [1, 2, 1, 2]), ("replace", [1, 2])]:
            data_tools.dispatch("result.export_table", {**params, "if_exists": mode}, namespace, store, connector)
            assert pd.read_sql(text('SELECT * FROM "temp"."shadow"'), connector.engine)["a"].tolist() == expected
        with pytest.raises(ValueError, match="already exists"):
            data_tools.dispatch("result.export_table", params, namespace, store, connector)
    finally:
        connector.disconnect()


def test_generic_pooled_connector_pins_same_engine_for_temporary_tables(tmp_path):
    class Connector:
        db_type = "sqlite"
        engine = create_engine("sqlite:///" + str(tmp_path / "generic.sqlite"), poolclass=QueuePool)
    connector = Connector()
    original = connector.engine
    from sqlalchemy import event
    checkouts = []
    @event.listens_for(original, "checkout")
    def namespace(driver, record, proxy):
        driver.execute("PRAGMA foreign_keys=ON")
        checkouts.append(driver)
    try:
        data_tools.dispatch("result.export_table", {"variable_name": "frame", "table": "temp_data", "temporary": True},
                            {"frame": pd.DataFrame({"a": [7]})}, ResultStore(pd, pl), connector)
        assert connector.engine is original and isinstance(original.pool, StaticPool)
        for _ in range(3):
            with original.connect() as connection:
                assert connection.execute(text("SELECT a FROM temp_data")).scalar() == 7
                assert connection.execute(text("PRAGMA foreign_keys")).scalar() == 1
        assert len(checkouts) >= 4 and len({id(driver) for driver in checkouts}) == 1
    finally:
        original.dispose()


def test_table_cancel_rolls_back_every_insert_batch():
    connector = SQLiteConnector(":memory:")
    try:
        connector.execute_query("CREATE TABLE destination(a INTEGER); INSERT INTO destination VALUES(99)")
        cancelled = [False]
        def update(item):
            if item["phase"] == "writing":
                cancelled[0] = True
        with pytest.raises(ExportCancelled):
            data_tools.dispatch("result.export_table", {"variable_name": "frame", "table": "destination", "if_exists": "append", "chunksize": 100},
                                {"frame": pd.DataFrame({"a": range(1000)})}, ResultStore(pd, pl), connector,
                                progress=update, cancelled=lambda: cancelled[0])
        assert connector.execute_query("SELECT a FROM destination")["a"].tolist() == [99]
    finally:
        connector.disconnect()


def test_retained_temporary_connection_is_not_reaped_or_evicted(monkeypatch):
    class Connector:
        has_temporary_tables = True
        def disconnect(self):
            raise AssertionError("Temporary session must not be evicted")
    pool = ConnectorPool(idle_timeout=1)
    for i in range(8):
        pool.items[(str(i), str(i))] = Connector()
    assert pool.reap_idle() == []
    with pytest.raises(ConnectionError, match="eight"):
        pool.activate({"config": {"db_type": "sqlite", "database": ":memory:"}})


@pytest.fixture
def client(tmp_path, monkeypatch):
    monkeypatch.setenv("DATAPYN_RUNTIME_STATE_PATH", str(tmp_path / "state"))
    monkeypatch.setenv("DATAPYN_WORKSPACE_PATH", str(tmp_path / "state"))
    value = Client()
    try:
        yield value
    finally:
        value.close()


def send(client, method, params):
    client.sequence += 1
    client.process.stdin.write(json.dumps({"id": client.sequence, "method": method, "params": params}) + "\n")
    client.process.stdin.flush()
    return client.sequence


def test_stdio_active_and_queued_cancel_preserve_kernel_variables(client, tmp_path):
    client.session()
    assert client.execute("marker = 73\nframe = pd.DataFrame({'value': range(200000)})", "seed")["status"] == "succeeded"
    request = send(client, "result.export", {"session_id": "a", "variable_name": "frame", "format": "xlsx",
                    "path": str(tmp_path / "cancel.xlsx"), "operation_id": "active-export"})
    client.wait(lambda message: message.get("event") == "result.export_progress"
                and message["payload"].get("operation_id") == "active-export")
    queued = send(client, "result.export", {"session_id": "a", "variable_name": "frame", "format": "csv",
                    "path": str(tmp_path / "queued.csv"), "operation_id": "queued-export"})
    assert client.request("result.export_cancel", {"session_id": "a", "operation_id": "queued-export"})["status"] == "cancelled"
    assert client.request("result.export_cancel", {"session_id": "a", "operation_id": "active-export"})["status"] == "cancelling"
    for identifier in (request, queued):
        result = client.wait(lambda message: message.get("id") == identifier)
        assert result["error"]["code"] == "cancelled"
    finished = client.execute("print(marker, len(frame))", "after-cancel")
    assert finished["status"] == "succeeded" and "73 200000" in client.output("after-cancel")
    assert not (tmp_path / "cancel.xlsx").exists() and not (tmp_path / "queued.csv").exists()
    assert not any(message.get("event") == "session.restarted" for message in client.all_messages)


def test_stdio_temporary_table_export_and_other_destination_restore_source(client, tmp_path):
    client.session()
    source, target = tmp_path / "source.sqlite", tmp_path / "target.sqlite"
    client.request("connection.connect", {"session_id": "a", "config": {"db_type": "sqlite", "database": str(source)}})
    seeded = client.execute("frame = pd.DataFrame({'id':[1,2]})\noriginal_engine = db_engine", "seed")
    assert seeded["status"] == "succeeded"
    client.request("result.export_table", {"session_id": "a", "variable_name": "frame", "table": "working", "temporary": True, "operation_id": "temp"})
    finished = client.execute("SELECT * FROM working", "read-temp", language="sql")
    assert client.page(finished)["rows"] == [[1], [2]]
    client.request("result.export_table", {"session_id": "a", "variable_name": "frame", "table": "copied", "schema": "main",
                   "config": {"db_type": "sqlite", "database": str(target)}, "operation_id": "other-target"})
    checked = client.execute("assert db_engine is original_engine\nassert db_database == " + repr(str(source)), "check-default")
    assert checked["status"] == "succeeded"
    checked = client.execute("SELECT * FROM working", "read-temp-again", language="sql")
    assert client.page(checked)["rows"] == [[1], [2]]
    with sqlite3.connect(target) as connection:
        assert connection.execute("SELECT * FROM copied").fetchall() == [(1,), (2,)]
    client.wait(lambda message: message.get("event") == "language.context_updated" and message["payload"].get("metadata_invalidated"))


def test_script_export_resolves_shared_parameters_with_custom_delimiter(tmp_path):
    path = tmp_path / "analysis.py"
    parameters = [{"id": "sharedparam:count", "name": "count", "sql_type": "integer", "value": "7", "enabled": True}]
    data_tools.export_script({"path": str(path), "shared_delimiter": "::name::", "shared_parameters": parameters,
                             "blocks": [{"language": "sql", "code": "SELECT ::count:: AS value", "name": "frame"},
                                        {"language": "python", "code": "count = ::count::"}]})
    content = path.read_text()
    compile(content, str(path), "exec")
    assert "shared_count" in content and "'shared_count': 7" in content and "count = 7" in content


def test_decimal_and_unsigned_sqlite_table_values_never_pass_through_float():
    connector = SQLiteConnector(":memory:")
    try:
        frame = pd.DataFrame({"decimal": [Decimal("1.234567890123456789")], "unsigned": pd.Series([2**64-1], dtype="uint64")})
        data_tools.dispatch("result.export_table", {"variable_name": "frame", "table": "data"}, {"frame": frame}, ResultStore(pd, pl), connector)
        assert connector.execute_query("SELECT * FROM data").values.tolist() == [["1.234567890123456789", str(2**64-1)]]
        assert frame.iloc[0, 0] == Decimal("1.234567890123456789") and frame["unsigned"].dtype == "uint64"
    finally:
        connector.disconnect()


def test_sql_ddl_text_fallback_emits_string_literals_and_separate_schema_is_literal():
    decimal = Decimal("1." + "2" * 50)
    frame = pd.DataFrame({"decimal": [decimal]})
    sql = export(frame, format="sql", options={"db_type": "sqlserver", "sql_mode": "create_insert"})["text"]
    assert "[decimal] NVARCHAR(MAX)" in sql and "N'" + str(decimal) + "'" in sql
    sql = export(frame, format="sql", options={"db_type": "postgresql", "sql_mode": "create_insert", "schema_name": "my.schema", "table_name": "data", "table_name_literal": True})["text"]
    assert '"my.schema"."data"' in sql and '"my"."schema"' not in sql
    sql = export(frame, format="sql", options={"db_type": "sqlite", "sql_mode": "create_insert"})["text"]
    with sqlite3.connect(":memory:") as connection:
        connection.executescript(sql)
        assert connection.execute("SELECT decimal FROM data").fetchone()[0] == str(decimal)


def test_sqlserver_temporary_exists_uses_bound_tempdb_name_and_ddl_compiles():
    from sqlalchemy import Column, Integer, MetaData, Table
    from sqlalchemy.dialects import mssql
    from sqlalchemy.schema import CreateTable
    from datapyn_runtime.table_export import _temporary_exists
    class Result:
        def scalar(self):
            return 42
    class Connection:
        def execute(self, statement, parameters):
            assert str(statement) == "SELECT OBJECT_ID(:name, 'U')"
            assert parameters == {"name": "tempdb..##a]b"}
            return Result()
    assert _temporary_exists(Connection(), "##a]b", "sqlserver")
    ddl = str(CreateTable(Table("##a]b", MetaData(), Column("value", Integer))).compile(dialect=mssql.dialect()))
    assert "CREATE TABLE [##a]]b]" in ddl


def test_temporary_tables_are_explorable_and_complete_without_qualified_catalog():
    from datapyn_runtime.explorer import ObjectExplorer
    connector = SQLiteConnector(":memory:")
    try:
        with pytest.raises(ValueError, match="single name"):
            data_tools.dispatch("result.export_table", {"variable_name": "frame", "table": 'work.table', "temporary": True},
                                {"frame": pd.DataFrame({"exact": [7]})}, ResultStore(pd, pl), connector)
        data_tools.dispatch("result.export_table", {"variable_name": "frame", "table": '"work.table"', "temporary": True},
                            {"frame": pd.DataFrame({"exact": [7]})}, ResultStore(pd, pl), connector)
        explorer = ObjectExplorer(connector)
        nodes = explorer.objects("main", "table")
        node = next(node for node in nodes if node.get("temporary"))
        assert node["schema"] == "temp"
        assert explorer.list({"node": node})["nodes"][0]["name"] == "exact"
        snapshot = explorer.completion_schema('SELECT * FROM temp."work.table"')
        assert any(table["key"] == "temp.work.table" for table in snapshot["tables"])
        assert snapshot["columns"]["temp.work.table"][0]["name"] == "exact"
    finally:
        connector.disconnect()


@pytest.mark.parametrize("orient", ["index", "columns"])
def test_json_rejects_stringified_index_collision(orient):
    with pytest.raises(ValueError, match="converted to strings"):
        export(pd.DataFrame({"x": [1, 2]}, index=[1, "1"]), format="json", options={"orient": orient})


@pytest.mark.parametrize("temporary", [False, True])
def test_sqlite_object_bigint_export_remains_exact(temporary):
    connector = SQLiteConnector(":memory:")
    try:
        data_tools.dispatch("result.export_table", {"variable_name": "frame", "table": "data", "temporary": temporary},
                            {"frame": pd.DataFrame({"value": pd.Series([2**100], dtype=object)})}, ResultStore(pd, pl), connector)
        assert connector.execute_query("SELECT value FROM data").iloc[0, 0] == str(2**100)
    finally:
        connector.disconnect()


def test_sqlserver_text_dtype_is_unicode_and_temporary_metadata_has_no_dbo_prefix():
    from datapyn_runtime.export_control import ExportControl
    from datapyn_runtime.table_export import _typed_frame
    from sqlalchemy import Column, MetaData, Table
    from sqlalchemy.dialects import mssql
    from sqlalchemy.schema import CreateTable
    from datapyn_runtime.sql_completion import RuntimeSqlAutoCompleteService
    frame, types = _typed_frame(pd.DataFrame({"value": ["ação", "漢字"]}), "sqlserver", ExportControl(2))
    ddl = str(CreateTable(Table("#sample", MetaData(), Column("value", types["value"]))).compile(dialect=mssql.dialect()))
    assert "NVARCHAR(max)" in ddl
    service = RuntimeSqlAutoCompleteService()
    service.set_schema({"db_type": "sqlserver", "tables": [{"name": "#sample", "key": "#sample", "schema": "", "temporary": True}],
                        "columns": {"#sample": [{"name": "value"}]}})
    entry = service._find_schema_entry("#sample")
    assert entry["detail"] == "#sample" and entry["schema"] == ""


def test_temp_metadata_shadows_unqualified_permanent_but_not_qualified_reference():
    from datapyn_runtime.sql_completion import RuntimeSqlAutoCompleteService
    service = RuntimeSqlAutoCompleteService()
    service.set_schema({"db_type": "sqlite", "current_schema": "main",
                        "tables": [{"name": "data", "key": "main.data", "schema": "main"},
                                   {"name": "data", "key": "temp.data", "schema": "temp", "temporary": True}],
                        "columns": {"main.data": [{"name": "permanent"}], "temp.data": [{"name": "working"}]}})
    assert service._find_schema_entry("data")["columns"] == [{"name": "working"}]
    assert service._find_schema_entry("data", "main")["columns"] == [{"name": "permanent"}]


def test_script_sql_and_notebook_resolve_custom_global_parameters(tmp_path):
    params = {"shared_parameters": [{"id": "sharedparam:text", "name": "text", "sql_type": "text", "value": "O'Reilly", "enabled": True}],
              "shared_delimiter": "::name::"}
    path = tmp_path / "analysis.sql"
    data_tools.export_script({**params, "path": str(path), "db_type": "sqlite",
                             "blocks": [{"language": "sql", "code": "SELECT ::text:: AS value, '::text::' AS literal"}]})
    with sqlite3.connect(":memory:") as connection:
        assert connection.execute(path.read_text()).fetchone() == ("O'Reilly", "::text::")
    path = tmp_path / "analysis.ipynb"
    data_tools.export_script({**params, "path": str(path), "blocks": [{"language": "python", "code": "value = ::text::"}]})
    code = "".join(json.loads(path.read_text())["cells"][0]["source"])
    namespace = {}
    exec(code, namespace)
    assert namespace["value"] == "O'Reilly"


def test_databricks_temp_export_uses_actual_dialect_types_without_forbidden_properties():
    from contextlib import contextmanager
    from databricks.sqlalchemy import DatabricksDialect
    statements, inserted = [], []
    class Rows:
        def mappings(self):
            return []
    class Connection:
        def in_transaction(self):
            return True
        def exec_driver_sql(self, statement):
            statements.append(statement)
            return Rows()
        def execute(self, statement, parameters):
            statements.append(str(statement.compile(dialect=DatabricksDialect())))
            inserted.extend(parameters)
    class Engine:
        dialect = DatabricksDialect()
        pool = StaticPool(creator=lambda: None)
        @contextmanager
        def begin(self):
            yield Connection()
    class Connector:
        db_type = "databricks"
        engine = Engine()
    connector = Connector()
    result = data_tools.dispatch("result.export_table", {"variable_name": "frame", "table": "sample", "temporary": True},
                                {"frame": pd.DataFrame({"id": [7], "text": ["ação"]})}, ResultStore(pd, pl), connector)
    ddl = next(statement for statement in statements if statement.startswith("CREATE TEMPORARY"))
    assert ddl == "CREATE TEMPORARY TABLE `sample` (`id` BIGINT, `text` STRING)"
    assert "USING" not in ddl and "TBLPROPERTIES" not in ddl
    assert inserted == [{"id": 7, "text": "ação"}] and result["row_count"] == 1


def test_stdio_drop_temp_refreshes_completion_and_releases_retained_connection(client):
    client.session()
    client.request("connection.connect", {"session_id": "a", "config": {"db_type": "sqlite", "database": ":memory:"}})
    assert client.execute("CREATE TABLE data(permanent INTEGER); INSERT INTO data VALUES(99)", "permanent", language="sql")["status"] == "succeeded"
    assert client.execute("frame = pd.DataFrame({'working':[7]})", "frame")["status"] == "succeeded"
    client.request("result.export_table", {"session_id": "a", "variable_name": "frame", "table": "data", "temporary": True})
    nodes = client.request("explorer.list", {"session_id": "a", "node": {"kind": "category", "category": "table", "schema": "main"}})["nodes"]
    assert any(table.get("temporary") for table in nodes)
    dropped = client.execute("DROP TABLE temp.data", "drop-temp", language="sql")
    assert dropped["status"] == "succeeded"
    checked = client.execute("SELECT * FROM data", "permanent-again", language="sql")
    assert client.page(checked)["columns"][0]["name"] == "permanent" and client.page(checked)["rows"] == [[99]]
    nodes = client.request("explorer.list", {"session_id": "a", "node": {"kind": "category", "category": "table", "schema": "main"}})["nodes"]
    assert not any(table.get("temporary") for table in nodes)


def test_dropped_temp_registry_is_removed_without_touching_permanent_shadow():
    from datapyn_runtime.table_export import refresh_temporary_tables
    from datapyn_runtime.explorer import ObjectExplorer
    connector = SQLiteConnector(":memory:")
    try:
        connector.execute_query("CREATE TABLE data(permanent INTEGER)")
        data_tools.dispatch("result.export_table", {"variable_name": "frame", "table": "data", "temporary": True},
                            {"frame": pd.DataFrame({"working": [7]})}, ResultStore(pd, pl), connector)
        assert connector.has_temporary_tables
        connector.execute_query("DROP TABLE temp.data")
        queries = []
        connector.connection.set_trace_callback(queries.append)
        refresh_temporary_tables(connector)
        assert not connector.has_temporary_tables and connector._datapyn_temporary_tables == {}
        assert len([query for query in queries if "sqlite_temp_master" in query]) == 1
        snapshot = ObjectExplorer(connector).completion_schema("SELECT * FROM data")
        assert {column["name"] for column in snapshot["columns"]["main.data"]} == {"permanent"}
    finally:
        connector.disconnect()
