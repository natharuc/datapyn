"""Query and Python result types survive the real stdio export route."""

from datetime import datetime

from openpyxl import load_workbook

from test_export_modes import client  # noqa: F401


def test_sql_casts_and_literal_identifiers_reach_excel_with_their_result_types(client, tmp_path):
    client.session()
    client.request("connection.connect", {"session_id": "a", "config": {"db_type": "sqlite"}})
    finished = client.execute("""
        SELECT CAST('0042' AS INTEGER) AS amount,
               CAST(123 AS TEXT) AS user_text,
               '00123456789012' AS identifier,
               CAST('2026-10-15' AS TEXT) AS user_date,
               '=1+1' AS literal
    """, "sql-types", language="sql")
    assert finished["status"] == "succeeded", finished
    # Paging serializes display values; exporting must still use the typed frame.
    assert client.page(finished)["rows"] == [[42, "123", "00123456789012", "2026-10-15", "=1+1"]]
    path = tmp_path / "query.xlsx"
    client.request("result.export", {"session_id": "a", "result_id": finished["results"][0]["result_id"],
                                    "path": str(path), "format": "xlsx"})
    book = load_workbook(path)
    try:
        cells = list(book.active[2])
        assert [cell.value for cell in cells] == [42, "123", "00123456789012", "2026-10-15", "=1+1"]
        assert [cell.data_type for cell in cells] == ["n", "s", "s", "s", "s"]
    finally:
        book.close()


def test_python_decimal_and_mixed_fillna_remain_typed_after_result_paging(client, tmp_path):
    client.session()
    finished = client.execute("""
from decimal import Decimal
from datetime import date
typed = pd.DataFrame({
    'amount': [Decimal('45.60'), None],
    'quantity': [7, None],
    'date': [date(2026, 10, 15), None],
    'identifier': ['00042', None],
}, dtype=object).fillna('N/A')
typed
""", "python-types")
    assert finished["status"] == "succeeded", finished
    assert client.page(finished)["rows"][0] == ["45.60", 7, "2026-10-15", "00042"]
    path = tmp_path / "python.xlsx"
    client.request("result.export", {"session_id": "a", "result_id": finished["results"][0]["result_id"],
                                    "path": str(path), "format": "xlsx"})
    book = load_workbook(path)
    try:
        cells = list(book.active[2])
        assert [cell.value for cell in cells] == [45.6, 7, datetime(2026, 10, 15), "00042"]
        assert [cell.data_type for cell in cells] == ["n", "n", "d", "s"]
        assert cells[0].number_format == "0.00"
        assert [cell.value for cell in book.active[3]] == ["N/A"] * 4
        assert [cell.data_type for cell in book.active[3]] == ["s"] * 4
    finally:
        book.close()
