"""XLSX cells retain query and Python scalar types instead of grid JSON types."""

from datetime import date, datetime, time, timedelta, timezone
from decimal import Decimal
from pathlib import Path
import sys

import numpy as np
from openpyxl import load_workbook
import pandas as pd
import polars as pl
import pytest

sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "source"))
from datapyn_runtime import data_tools
from datapyn_runtime.kernel import ResultStore


def workbook(tmp_path, frame, *, by_result=False, options=None):
    store = ResultStore(pd, pl)
    reference = store.register(frame, "data")
    params = {"result_id": reference["result_id"]} if by_result else {"variable_name": "data"}
    path = tmp_path / "typed.xlsx"
    data_tools.dispatch("result.export", {**params, "path": str(path), "options": options or {}}, {"data": frame}, store)
    return load_workbook(path)


@pytest.mark.parametrize("by_result", [False, True])
def test_sql_decimal_values_and_explicit_text_have_distinct_excel_types(tmp_path, by_result):
    # SQL Server DECIMAL arrives as Decimal. CAST(... AS VARCHAR) arrives as str.
    frame = pd.DataFrame({
        "amount": [Decimal("1234.50"), Decimal("0.01")],
        "integer": pd.Series([17, None], dtype="Int64"),
        "flag": pd.Series([True, None], dtype="boolean"),
        "cast_as_text": ["001234.50", "33608308000173"],
        "date_as_text": ["2026-10-15", "01"],
    })
    book = workbook(tmp_path, frame, by_result=by_result)
    try:
        sheet = book.active
        assert (sheet["A2"].value, sheet["A2"].data_type) == (1234.5, "n")
        assert (sheet["A3"].value, sheet["A3"].data_type) == (0.01, "n")
        assert sheet["A2"].number_format == "0.00"
        assert (sheet["B2"].value, sheet["B2"].data_type) == (17, "n")
        assert sheet["B3"].value is None
        assert (sheet["C2"].value, sheet["C2"].data_type) == (True, "b")
        assert sheet["C3"].value is None
        assert [(sheet[f"D{row}"].value, sheet[f"D{row}"].data_type) for row in (2, 3)] == [
            ("001234.50", "s"), ("33608308000173", "s"),
        ]
        assert (sheet["E2"].value, sheet["E2"].data_type) == ("2026-10-15", "s")
        assert (sheet["E3"].value, sheet["E3"].data_type) == ("01", "s")
    finally:
        book.close()


def test_fillna_object_columns_preserve_each_values_type(tmp_path):
    frame = pd.DataFrame({
        "amount": [Decimal("25.75"), None],
        "date": pd.Series([date(2026, 10, 15), None], dtype=object),
        "when": pd.Series([datetime(2026, 10, 15, 12, 34, 56, 123000), None], dtype=object),
        "time": [time(12, 34, 56), None],
    }).fillna("N/A")
    book = workbook(tmp_path, frame)
    try:
        sheet = book.active
        assert (sheet["A2"].value, sheet["A2"].data_type) == (25.75, "n")
        assert sheet["B2"].value == datetime(2026, 10, 15)
        assert sheet["B2"].data_type == "d" and sheet["B2"].number_format == "yyyy-mm-dd"
        assert sheet["C2"].value == datetime(2026, 10, 15, 12, 34, 56, 123000)
        assert sheet["C2"].data_type == "d"
        assert (sheet["D2"].value, sheet["D2"].data_type) == (time(12, 34, 56), "d")
        assert [(cell.value, cell.data_type) for cell in sheet[3]] == [("N/A", "s")] * 4
    finally:
        book.close()


def test_excel_precision_limit_does_not_silently_round_exact_values(tmp_path):
    frame = pd.DataFrame({"value": pd.Series([
        999999999999999, 1000000000000001, 2**60,
        Decimal("1234567890123.45"), Decimal("1234567890123.456"),
        Decimal("0.00000000000000000001"), Decimal("1E-400"), Decimal("1E400"),
    ], dtype=object)})
    book = workbook(tmp_path, frame)
    try:
        cells = list(book.active.iter_rows(min_row=2, max_col=1))
        assert [(row[0].value, row[0].data_type) for row in cells] == [
            (999999999999999, "n"), ("1000000000000001", "s"), (str(2**60), "s"),
            (1234567890123.45, "n"), ("1234567890123.456", "s"),
            (1e-20, "n"), ("1E-400", "s"), ("1E+400", "s"),
        ]
    finally:
        book.close()


def test_numpy_scalars_missing_values_formulas_and_timezone_are_not_reinterpreted(tmp_path):
    aware = datetime(2026, 10, 15, 12, 0, tzinfo=timezone(timedelta(hours=-3)))
    frame = pd.DataFrame({"value": pd.Series([
        np.int64(42), np.float32(1.5), np.bool_(True), Decimal("NaN"),
        pd.NA, pd.NaT, np.float64("nan"), np.float32("inf"),
        "=SUM(A1:A3)", "+42", "@name", aware,
    ], dtype=object)})
    book = workbook(tmp_path, frame)
    try:
        sheet = book.active
        assert [(sheet[f"A{row}"].value, sheet[f"A{row}"].data_type) for row in (2, 3, 4)] == [
            (42, "n"), (1.5, "n"), (True, "b"),
        ]
        assert all(sheet[f"A{row}"].value is None for row in range(5, 10))
        assert [(sheet[f"A{row}"].value, sheet[f"A{row}"].data_type) for row in range(10, 14)] == [
            ("=SUM(A1:A3)", "s"), ("+42", "s"), ("@name", "s"), (aware.isoformat(), "s"),
        ]
    finally:
        book.close()


def test_polars_decimal_export_uses_the_native_result_values(tmp_path):
    frame = pl.DataFrame({
        "amount": pl.Series([Decimal("12.50"), None], dtype=pl.Decimal(18, 2)),
        "code": ["00012", "00336"],
    })
    book = workbook(tmp_path, frame, by_result=True)
    try:
        sheet = book.active
        assert (sheet["A2"].value, sheet["A2"].data_type) == (12.5, "n")
        assert sheet["A3"].value is None
        assert (sheet["B2"].value, sheet["B2"].data_type) == ("00012", "s")
    finally:
        book.close()


def test_native_timestamps_and_durations_keep_excel_temporal_cells(tmp_path):
    elapsed = timedelta(days=2, hours=3, minutes=4, seconds=5, milliseconds=123)
    frame = pd.DataFrame({"value": pd.Series([
        np.datetime64("2026-10-15T12:34:56.123", "ns"),
        np.timedelta64(51, "h"), elapsed, pd.Timedelta(elapsed),
        time(12, 34, 56, 123000),
        time(12, 34, tzinfo=timezone(timedelta(hours=-3))),
    ], dtype=object)})
    book = workbook(tmp_path, frame)
    try:
        sheet = book.active
        assert (sheet["A2"].value, sheet["A2"].data_type) == (datetime(2026, 10, 15, 12, 34, 56, 123000), "d")
        assert sheet["A3"].value == timedelta(hours=51)
        assert sheet["A3"].number_format == "[h]:mm:ss"
        assert sheet["A4"].value == sheet["A5"].value == elapsed
        assert sheet["A4"].number_format == "[h]:mm:ss.000"
        assert sheet["A6"].value == time(12, 34, 56, 123000)
        assert sheet["A6"].number_format == "hh:mm:ss.000"
        assert (sheet["A7"].value, sheet["A7"].data_type) == ("12:34:00-03:00", "s")
    finally:
        book.close()


def test_temporal_values_outside_excel_precision_and_date_range_remain_exact(tmp_path):
    values = [
        datetime(2026, 10, 15, 12, 34, 56, 123456),
        time(12, 34, 56, 123456),
        pd.Timestamp("2026-10-15T12:34:56.000000001"),
        np.datetime64("2026-10-15T12:34:56.123456789", "ns"),
        date(1899, 12, 31), datetime(1899, 12, 31, 23, 59, 59),
        date(1753, 1, 1),
        timedelta(microseconds=123456), pd.Timedelta(nanoseconds=1),
        np.timedelta64(123456789, "ns"),
        timedelta(hours=-1), pd.Timedelta(days=-2),
        np.datetime64("10000-01-01T00:00:00", "s"),
    ]
    frame = pd.DataFrame({"value": pd.Series(values + [None], dtype=object)}).fillna("N/A")
    book = workbook(tmp_path, frame)
    try:
        cells = [row[0] for row in book.active.iter_rows(min_row=2)]
        assert [(cell.value, cell.data_type) for cell in cells] == [
            ("2026-10-15T12:34:56.123456", "s"),
            ("12:34:56.123456", "s"),
            ("2026-10-15T12:34:56.000000001", "s"),
            ("2026-10-15T12:34:56.123456789", "s"),
            ("1899-12-31", "s"), ("1899-12-31T23:59:59", "s"),
            ("1753-01-01", "s"),
            ("0:00:00.123456", "s"), ("0 days 00:00:00.000000001", "s"),
            ("0 days 00:00:00.123456789", "s"),
            ("-1 day, 23:00:00", "s"), ("-2 days +00:00:00", "s"),
            ("10000-01-01T00:00:00", "s"), ("N/A", "s"),
        ]
    finally:
        book.close()
